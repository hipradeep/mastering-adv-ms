from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_lld():
    document = Document()

    # Title
    document.add_heading('Low Level Design (LLD)', 0)
    p = document.add_paragraph('Terminator Project')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Database Design
    document.add_heading('1. Database Design (Schema)', level=1)
    document.add_paragraph('The system is built on a Service-per-DB pattern using PostgreSQL. Below are the schema definitions for the core tables.')

    # 1.1 Orchestrator
    document.add_heading('1.1 Orchestrator Service', level=2)
    document.add_paragraph('Table: hstt_saga_transactions (Schema: saga)')
    document.add_paragraph('Tracks the lifecycle of distributed transactions.')
    
    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column Name'
    hdr_cells[1].text = 'Data Type'
    hdr_cells[2].text = 'Constraints'
    hdr_cells[3].text = 'Description'

    orch_rows = [
        ('hststr_transaction_id', 'VARCHAR', 'PK', 'Unique UUID for the Saga.'),
        ('hststr_status', 'VARCHAR', 'Enum', 'STARTED, SUCCESS, FAILURE.'),
        ('hststr_request_payload', 'TEXT', '', 'JSON dump of the original IssueRequestDto.'),
        ('hststr_failure_reason', 'VARCHAR', '', 'Error message if failed.'),
        ('gnum_hospital_code', 'INTEGER', '', 'Context identifier (e.g., 998).'),
        ('hstdt_completion_date', 'TIMESTAMP', '', 'Time when saga ended.'),
        ('gdt_entry_date', 'TIMESTAMP', '', 'Audit: Creation time.')
    ]
    for row in orch_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = row[0]
        row_cells[1].text = row[1]
        row_cells[2].text = row[2]
        row_cells[3].text = row[3]

    # 1.2 Inventory
    document.add_heading('1.2 Inventory Service', level=2)
    document.add_paragraph('Table: hstt_drug_currstock_dtl (Schema: dwh)')
    document.add_paragraph('Holds the real-time stock position for items in stores.')

    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column Name'
    hdr_cells[1].text = 'Data Type'
    hdr_cells[2].text = 'Constraints'
    hdr_cells[3].text = 'Description'

    inv_rows = [
        ('hstnum_store_id', 'INTEGER', 'PK', 'Store Identifier.'),
        ('hstnum_itembrand_id', 'INTEGER', 'PK', 'Drug/Item Identifier.'),
        ('hststr_batch_no', 'VARCHAR', 'PK', 'Specific Batch Number.'),
        ('hstnum_mfg_id', 'INTEGER', 'PK', 'Manufacturer ID.'),
        ('hstnum_programme_id', 'INTEGER', 'PK', 'Programme/Scheme ID.'),
        ('gnum_hospital_code', 'INTEGER', 'PK', 'Hospital Code.'),
        ('hstnum_stock_status_code', 'INTEGER', 'PK', 'Status (e.g., 10=Active).'),
        ('hstnum_inhand_qty', 'DOUBLE', '', 'Critical: Current usable quantity.'),
        ('hstnum_reserved_qty', 'DOUBLE', '', 'Quantity locked by ongoing Sagas.'),
        ('hstdt_expiry_date', 'DATE', '', 'Batch expiry date.'),
        ('hstnum_rate', 'DOUBLE', '', 'Purchase rate per unit.')
    ]
    for row in inv_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = row[0]
        row_cells[1].text = row[1]
        row_cells[2].text = row[2]
        row_cells[3].text = row[3]

    # 1.3 Issue Service
    document.add_heading('1.3 Issue Service', level=2)
    document.add_paragraph('Table: hstt_patemp_issue_dtl')
    document.add_paragraph('Header table for a patient issue transaction.')

    table = document.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Column Name'
    hdr_cells[1].text = 'Data Type'
    hdr_cells[2].text = 'Constraints'
    hdr_cells[3].text = 'Description'

    issue_rows = [
        ('hstnum_issue_no', 'INTEGER', 'PK', 'Unique Issue Number.'),
        ('hstnum_store_id', 'INTEGER', 'PK', 'Issuing Store.'),
        ('gnum_hospital_code', 'INTEGER', 'PK', 'Hospital Code.'),
        ('hrgnum_puk', 'VARCHAR', '', 'Patient Unique Key (CR No).'),
        ('hstdt_issue_date', 'DATE', '', 'Date of issue.'),
        ('hstnum_net_cost', 'DOUBLE', '', 'Total cost of issue.')
    ]
    for row in issue_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = row[0]
        row_cells[1].text = row[1]
        row_cells[2].text = row[2]
        row_cells[3].text = row[3]


    # 2. API Specifications
    document.add_heading('2. API Specifications', level=1)
    
    document.add_heading('2.1 Orchestrator Service', level=2)
    document.add_paragraph('Endpoint: POST /api/orchestrator/issue').bold = True
    document.add_paragraph('Initiates the Issue Drug Saga.')
    document.add_paragraph('Request Body (IssueRequestDto):')
    document.add_paragraph(
        '{\n'
        '  "gnumHospitalCode": 998,\n'
        '  "hstnumStoreId": 101,\n'
        '  "hrgnumPuk": "1000001",\n'
        '  "issueItems": [\n'
        '    {\n'
        '      "hstnumItemBrandId": 5001,\n'
        '      "hststrBatchNo": "B001",\n'
        '      "hstnumIssueQty": 10\n'
        '    }\n'
        '  ]\n'
        '}'
    )

    document.add_heading('2.2 Inventory Service', level=2)
    document.add_paragraph('Endpoint: GET /api/inventory/stock').bold = True
    document.add_paragraph('Fetches available stock for the UI.')
    document.add_paragraph('Parameters: hospitalCode (default 998), storeId, itemBrandId.')

    # 3. Implementation Logic
    document.add_heading('3. Implementation Logic Details', level=1)
    
    document.add_heading('3.1 Atomic Inventory Updates', level=2)
    document.add_paragraph(
        'To prevent Race Conditions (Overselling) without using performance-heavy database locks (Pessimistic Locking), '
        'the system uses **Atomic Native SQL Queries** in DrugCurrstockDtlRepository.'
    )
    document.add_paragraph('Logic:')
    document.add_paragraph(
        'UPDATE dwh.hstt_drug_currstock_dtl \n'
        'SET hstnum_inhand_qty = hstnum_inhand_qty - :deductQty \n'
        'WHERE hstnum_store_id = :storeId \n'
        '  AND hstnum_itembrand_id = :itemBrandId \n'
        '  AND hststr_batch_no = :batchNo \n'
        '  AND hstnum_inhand_qty >= :deductQty'
    )
    document.add_paragraph(
        'Explanation: The condition "hstnum_inhand_qty >= :deductQty" ensures that if two users try to deduct the last '
        'units simultaneously, the database will only allow the first one to succeed (Row Count = 1). The second one will '
        'fail (Row Count = 0).'
    )

    document.add_heading('3.2 Saga Compensation', level=2)
    document.add_paragraph('If the Issue Service fails to save the record (e.g., DB Constraint Violation):')
    
    steps = [
        'Issue Service publishes "IssueCreatedEvent" with success=false.',
        'Orchestrator receives failure event.',
        'Orchestrator sends "ConfirmStockCommand" with commit=false.',
        'Inventory Service receives command and Rolls Back the stock deductions.',
        'Saga Status updated to FAILURE.'
    ]
    for step in steps:
        document.add_paragraph(step, style='List Number')

    # Save
    output_path = 'd:/Projects Workspaces/terminator/md_files/Low_Level_Design.docx'
    document.save(output_path)
    print(f"Document saved to {output_path}")

if __name__ == '__main__':
    create_lld()
